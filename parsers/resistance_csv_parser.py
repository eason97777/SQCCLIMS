#!/usr/bin/env python3
import csv
import math
import re
import statistics
from pathlib import Path


PARSER_NAME = "resistance_csv_parser"
PARSER_VERSION = "0.4.0"
SCHEMA_VERSION = "1.0"

MATRIX_SIZE = 12
BLOCK_SIZE = MATRIX_SIZE + 1
BLANK_ROWS_BETWEEN_BLOCKS = 1
FULL_DIE_LAYOUT = [
    [None, None, None, None, "A1", "A2", "A3", "A4", "A5", None, None, None, None],
    [None, None, None, "B1", "B2", "B3", "B4", "B5", "B6", "B7", None, None, None],
    [None, None, "C1", "C2", "C3", "C4", "C5", "C6", "C7", "C8", "C9", None, None],
    [None, "D1", "D2", "D3", "D4", "D5", "D6", "D7", "D8", "D9", "D10", "D11", None],
    ["E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8", "E9", "E10", "E11", "E12", "E13"],
    ["F1", "F2", "F3", "F4", "F5", "F6", "F7", "F8", "F9", "F10", "F11", "F12", "F13"],
    ["G1", "G2", "G3", "G4", "G5", "G6", "G7", "G8", "G9", "G10", "G11", "G12", "G13"],
    ["H1", "H2", "H3", "H4", "H5", "H6", "H7", "H8", "H9", "H10", "H11", "H12", "H13"],
    ["I1", "I2", "I3", "I4", "I5", "I6", "I7", "I8", "I9", "I10", "I11", "I12", "I13"],
    [None, "J1", "J2", "J3", "J4", "J5", "J6", "J7", "J8", "J9", "J10", "J11", None],
    [None, None, "K1", "K2", "K3", "K4", "K5", "K6", "K7", "K8", "K9", None, None],
    [None, None, None, "L1", "L2", "L3", "L4", "L5", "L6", "L7", None, None, None],
    [None, None, None, None, "M1", "M2", "M3", "M4", "M5", None, None, None, None],
]
FULL_DIE_IDS = {
    die_id
    for row in FULL_DIE_LAYOUT
    for die_id in row
    if die_id is not None
}
NA_VALUES = {"", "na", "n/a", "nan", "null", "none", "-", "--"}


def clean_cell(value):
    return str(value or "").replace("\ufeff", "").strip()


def trim_trailing_empty_cells(row):
    trimmed = list(row)
    while trimmed and clean_cell(trimmed[-1]) == "":
        trimmed.pop()
    return trimmed


def is_blank_row(row):
    return all(clean_cell(value) == "" for value in row)


def normalize_number_text(value):
    return (
        clean_cell(value)
        .replace(",", "")
        .replace("，", "")
        .replace("．", ".")
        .replace("－", "-")
    )


def parse_die_id(value, row_number):
    die_id = clean_cell(value).upper()
    if not re.fullmatch(r"[A-M](?:[1-9]|1[0-3])", die_id):
        raise ValueError(f"第 {row_number} 行 Die 编号无效：{value}")
    return die_id


def parse_resistance_value(value, row_number, col_number):
    text = normalize_number_text(value)
    if text.lower() in NA_VALUES:
        return text, None
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"第 {row_number} 行第 {col_number} 列电阻值不是有效数字：{value}") from exc
    if not math.isfinite(number):
        raise ValueError(f"第 {row_number} 行第 {col_number} 列电阻值不是有限数字：{value}")
    return text, number


def point_area(row_index, col_index):
    if row_index <= 6 and col_index <= 6:
        return "A"
    if row_index <= 6 and col_index > 6:
        return "B"
    if row_index > 6 and col_index <= 6:
        return "C"
    return "D"


def format_number(value, digits=4):
    if value is None:
        return None
    rounded = round(value, digits)
    return int(rounded) if float(rounded).is_integer() else rounded


def summarize_values(records, *, die_id=None, area=None):
    total = len(records)
    valid_values = [
        record["numeric_value"]
        for record in records
        if isinstance(record.get("numeric_value"), (int, float))
    ]
    count_valid = len(valid_values)
    count_outlier = sum(1 for record in records if record.get("is_outlier"))
    count_na = total - count_valid
    count_normal = count_valid - count_outlier
    average = statistics.mean(valid_values) if valid_values else None
    maximum = max(valid_values) if valid_values else None
    minimum = min(valid_values) if valid_values else None
    value_range = maximum - minimum if maximum is not None and minimum is not None else None
    uniformity = (
        value_range / (2 * average)
        if value_range is not None and average not in (None, 0)
        else None
    )
    summary = {
        "count_total": total,
        "count_valid": count_valid,
        "count_na": count_na,
        "count_outlier": count_outlier,
        "count_normal": count_normal,
        "yield_rate": format_number(count_normal / total if total else None),
        "max": format_number(maximum),
        "min": format_number(minimum),
        "range": format_number(value_range),
        "average": format_number(average),
        "uniformity": format_number(uniformity),
    }
    if die_id is not None:
        summary["die_id"] = die_id
    if area is not None:
        summary["area"] = area
    return summary


def read_csv_rows(path):
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [trim_trailing_empty_cells(row) for row in csv.reader(handle)]
    except UnicodeDecodeError as exc:
        raise ValueError("CSV 编码不是 UTF-8，请将 CSV 转换为 UTF-8 后再导入") from exc


def read_xlsx_rows(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Resistance parser supports XLSX only when openpyxl is installed; please convert the file to CSV.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(trim_trailing_empty_cells(["" if value is None else value for value in row]))
        return rows
    finally:
        workbook.close()


def read_table_rows(path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix == ".xlsx":
        return read_xlsx_rows(path)
    if suffix == ".xls":
        raise ValueError("Resistance parser does not support XLS in this environment; please convert the file to XLSX or CSV.")
    raise ValueError("Resistance parser supports CSV and XLSX files.")


def next_non_empty_row(rows, start_index):
    index = start_index
    while index < len(rows) and is_blank_row(rows[index]):
        index += 1
    return index


def parse_die_block(rows, start_index, context, record_index_start):
    if start_index + BLOCK_SIZE > len(rows):
        raise ValueError(f"第 {start_index + 1} 行开始的 Die block 不完整，需要 13 行")

    header_row = rows[start_index]
    if len(header_row) < BLOCK_SIZE:
        raise ValueError(f"第 {start_index + 1} 行表头不完整，需要 Die 编号和 12 个列表头")

    die_id = parse_die_id(header_row[0], start_index + 1)
    col_headers = [clean_cell(value) for value in header_row[1:BLOCK_SIZE]]
    if any(header == "" for header in col_headers):
        raise ValueError(f"Die {die_id} 的列表头不能为空")

    records = []
    for row_offset in range(MATRIX_SIZE):
        source_row_index = start_index + 1 + row_offset
        source_row = rows[source_row_index]
        if len(source_row) < BLOCK_SIZE:
            raise ValueError(f"Die {die_id} 第 {row_offset + 1} 个数据行不完整，需要 13 列")

        row_header = clean_cell(source_row[0])
        if not row_header:
            raise ValueError(f"Die {die_id} 第 {row_offset + 1} 个行表头不能为空")

        for col_offset in range(MATRIX_SIZE):
            cleaned_value, numeric_value = parse_resistance_value(
                source_row[col_offset + 1],
                source_row_index + 1,
                col_offset + 2,
            )
            row_index = row_offset + 1
            col_index = col_offset + 1
            record_index = record_index_start + len(records)
            record = {
                **context,
                "record_index": record_index,
                "die_id": die_id,
                "row_index": row_index,
                "col_index": col_index,
                "row_header": row_header,
                "col_header": col_headers[col_offset],
                "area": point_area(row_index, col_index),
                "primary_key": f"{die_id}:R{row_index:02d}:C{col_index:02d}",
                "raw_value": clean_cell(source_row[col_offset + 1]),
                "cleaned_value": cleaned_value,
                "numeric_value": numeric_value,
                "is_outlier": False,
                "outlier_reason": None,
            }
            records.append(record)

    return die_id, records


def build_summaries(records, die_order):
    die_summary = []
    area_summary = []
    for die_id in die_order:
        die_records = [record for record in records if record["die_id"] == die_id]
        die_summary.append(summarize_values(die_records, die_id=die_id))
        for area in ("A", "B", "C", "D"):
            area_records = [record for record in die_records if record["area"] == area]
            area_summary.append(summarize_values(area_records, die_id=die_id, area=area))
    return summarize_values(records), die_summary, area_summary


def parse_resistance_csv(
    csv_path,
    *,
    raw_data_id,
    sample_id,
    sample_uid,
    raw_data_code,
    sample_display_code="",
    data_type="resistance",
    source_file_info=None,
):
    path = Path(csv_path)
    rows = read_table_rows(path)
    warnings = []
    records = []
    die_order = []
    seen_die_ids = set()
    row_index = next_non_empty_row(rows, 0)
    source_file = path.name
    context = {
        "raw_data_id": raw_data_id,
        "sample_id": sample_id,
        "sample_uid": sample_uid,
        "raw_data_code": raw_data_code,
        "sample_display_code": sample_display_code,
        "data_type": data_type,
        "source_file": source_file,
    }

    while row_index < len(rows):
        die_id, block_records = parse_die_block(rows, row_index, context, len(records) + 1)
        if die_id in seen_die_ids:
            raise ValueError(f"Die {die_id} 重复出现")
        seen_die_ids.add(die_id)
        die_order.append(die_id)
        records.extend(block_records)

        row_index += BLOCK_SIZE
        blank_start = row_index
        while row_index < len(rows) and is_blank_row(rows[row_index]):
            row_index += 1
        blank_count = row_index - blank_start
        if row_index < len(rows) and blank_count != BLANK_ROWS_BETWEEN_BLOCKS:
            raise ValueError(
                f"Die {die_id} 后应固定空 1 行，实际为空 {blank_count} 行"
            )

    if not records:
        raise ValueError("CSV 中没有有效 Resistance Die block")

    unknown_die_ids = [die_id for die_id in die_order if die_id not in FULL_DIE_IDS]
    if unknown_die_ids:
        raise ValueError(
            "CSV contains Die IDs outside the full wafer layout: "
            + ", ".join(unknown_die_ids)
        )

    missing_die_ids = sorted(FULL_DIE_IDS - set(die_order))

    overall_summary, die_summary, area_summary = build_summaries(records, die_order)
    summary = {
        "die_layout": "full_wafer_die",
        "die_count": len(die_order),
        "expected_die_count": len(FULL_DIE_IDS),
        "die_ids": die_order,
        "missing_die_ids": missing_die_ids,
        "total_point_count": len(records),
        "count_valid": overall_summary["count_valid"],
        "count_na": overall_summary["count_na"],
        "count_outlier": overall_summary["count_outlier"],
        "count_normal": overall_summary["count_normal"],
        "yield_rate": overall_summary["yield_rate"],
        "overall_summary": overall_summary,
        "die_summary": die_summary,
        "area_summary": area_summary,
        "schema": {
            "parser_mode": "die_matrix",
            "matrix_size": "12x12",
            "block_size": "13x13",
            "blank_rows_between_blocks": BLANK_ROWS_BETWEEN_BLOCKS,
            "die_id_cell": "top_left",
            "col_header_row": 1,
            "row_header_col": 1,
            "yield_rate_formula": "count_normal / count_total",
            "cleaning_config": {
                "enabled": False,
                "lower_limit": None,
                "upper_limit": None,
            },
        },
        "source_file": source_file_info or {"original_filename": source_file},
    }

    return {
        "schema_version": SCHEMA_VERSION,
        "data_type": data_type,
        "raw_data_id": raw_data_id,
        "sample_id": sample_id,
        "sample_uid": sample_uid,
        "sample_display_code": sample_display_code,
        "raw_data_code": raw_data_code,
        "parser_name": PARSER_NAME,
        "parser_version": PARSER_VERSION,
        "records": records,
        "summary": summary,
        "plots": [],
        "warnings": warnings,
        "errors": [],
    }
